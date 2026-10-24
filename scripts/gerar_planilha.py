#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera dados/planilha_pib.xlsx — a mesma base tratada, em formato amigável ao Excel.

Por que existe: os CSV do projeto usam ponto decimal (formato neutro, ideal para o
Power BI). Ao abrir um CSV com duplo clique no Excel em português, os números podem
aparecer como texto. A planilha .xlsx resolve isso: é binária, não depende de
delimitador e já traz os números com formato brasileiro (R$, %, índices).

Uso (a partir da raiz do projeto):
    python scripts/gerar_planilha.py
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
DATA_ACESSO = "2026-10-20"

COR_CABECALHO = "0F4C5C"
COR_ZEBRA = "F6F9FA"
FONTE = "Segoe UI"


def numero_br(valor: float, casas: int = 2) -> str:
    texto = f"{valor:,.{casas}f}"
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


def brl_milhoes(valor_mi: float) -> str:
    reais = valor_mi * 1_000_000
    for escala, singular, plural in (
        (1e12, "trilhão", "trilhões"),
        (1e9, "bilhão", "bilhões"),
        (1e6, "milhão", "milhões"),
    ):
        if abs(reais) >= escala:
            valor = reais / escala
            nome = singular if abs(valor) < 2 else plural
            return f"R$ {numero_br(valor)} {nome}"
    return f"R$ {numero_br(reais, 0)}"


FORMATOS = {
    "Milhões de Reais": 'R$ #,##0',
    "Número-índice": "0.00",
    "%": '0.0"%"',
}


def estilo_cabecalho(ws, largura_linha: int) -> None:
    fill = PatternFill("solid", fgColor=COR_CABECALHO)
    fonte = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
    for coluna in range(1, largura_linha + 1):
        celula = ws.cell(row=1, column=coluna)
        celula.fill = fill
        celula.font = fonte
        celula.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def ajustar_larguras(ws) -> None:
    for coluna in ws.columns:
        maior = max(len(str(c.value)) if c.value is not None else 0 for c in coluna)
        ws.column_dimensions[get_column_letter(coluna[0].column)].width = min(maior + 3, 60)


def faixas(ws, linhas: int, colunas: int) -> None:
    borda = Border(bottom=Side(style="thin", color="E3E8EC"))
    for linha in range(2, linhas + 2):
        for coluna in range(1, colunas + 1):
            celula = ws.cell(row=linha, column=coluna)
            celula.font = Font(name=FONTE, size=10)
            celula.border = borda
            if linha % 2 == 0:
                celula.fill = PatternFill("solid", fgColor=COR_ZEBRA)


def aba_leia_me(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Leia-me"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 86

    linhas = [
        ("Projeto", "BI — Tema 12: PIB Brasileiro"),
        ("Aluno", "Evandro Miguel · Turma 1"),
        ("Pergunta norteadora",
         "Como o PIB brasileiro evoluiu e como os principais setores participaram desse desempenho?"),
        ("Fonte", "IBGE — SIDRA, tabelas 1620, 1621, 1846 e 5932 (Contas Nacionais Trimestrais)"),
        ("Data de acesso", DATA_ACESSO),
        ("Período", "1º trimestre de 2000 a 2º trimestre de 2026 (106 trimestres)"),
        ("", ""),
        ("Aba Fatos", "Uma linha por tabela x série x trimestre. Coluna 'unidade' indica a grandeza."),
        ("Aba dim_tempo", "Dimensão de tempo trimestral (data, ano, trimestre, rótulo)."),
        ("Aba dim_serie", "Dimensão das séries (grupo, subgrupo, nível e ordem de exibição)."),
        ("Aba Destaques", "Indicadores principais já calculados, para leitura rápida."),
        ("", ""),
        ("Cuidado 1", "2026 é parcial: existem apenas o 1º e o 2º trimestres."),
        ("Cuidado 2", "Valores correntes (R$) não são comparáveis em nível entre anos; use o índice de volume."),
        ("Cuidado 3", "A série 1621 tem ajuste sazonal: serve para comparação trimestre contra trimestre."),
        ("Cuidado 4", "Importação vem com sinal positivo e deve ser subtraída ao compor a demanda."),
        ("Cuidado 5", "Índices e percentuais não devem ser somados; somas só fazem sentido em R$ correntes."),
    ]

    ws["A1"] = "Base tratada do PIB brasileiro"
    ws["A1"].font = Font(name=FONTE, bold=True, size=15, color=COR_CABECALHO)
    ws.merge_cells("A1:B1")

    for i, (rotulo, valor) in enumerate(linhas, start=3):
        ws.cell(row=i, column=1, value=rotulo).font = Font(name=FONTE, bold=True, size=10)
        ws.cell(row=i, column=2, value=valor).font = Font(name=FONTE, size=10)
        ws.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")


def aba_fatos(wb: Workbook, dados: dict[str, pd.DataFrame]) -> None:
    ws = wb.create_sheet("Fatos")
    colunas = ["tabela", "trimestre_txt", "serie", "valor", "unidade", "ano", "trimestre", "data"]
    ws.append(["Tabela", "Trimestre", "Série", "Valor", "Unidade", "Ano", "Trim.", "Data"])

    for tabela, df in dados.items():
        for _, r in df.iterrows():
            ws.append([
                tabela,
                str(r.trimestre_txt),
                str(r.serie),
                float(r.valor),
                str(r.unidade),
                int(r.ano),
                int(r.trimestre),
                str(r.data),
            ])
            linha = ws.max_row
            formato = FORMATOS.get(str(r.unidade))
            if formato:
                ws.cell(row=linha, column=4).number_format = formato

    for linha in range(2, ws.max_row + 1):
        ws.cell(row=linha, column=4).alignment = Alignment(horizontal="right")
        ws.cell(row=linha, column=6).alignment = Alignment(horizontal="center")
        ws.cell(row=linha, column=7).alignment = Alignment(horizontal="center")

    estilo_cabecalho(ws, len(colunas))
    ajustar_larguras(ws)
    ws.column_dimensions["C"].width = 62
    faixas(ws, ws.max_row - 1, len(colunas))


def aba_dimensao(wb: Workbook, nome: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(nome)
    ws.append(list(df.columns))
    for _, r in df.iterrows():
        linha = []
        for coluna in df.columns:
            valor = r[coluna]
            if pd.isna(valor):
                linha.append(None)
            elif pd.api.types.is_integer_dtype(df[coluna]):
                linha.append(int(valor))
            else:
                linha.append(str(valor))
        ws.append(linha)
    estilo_cabecalho(ws, len(df.columns))
    ajustar_larguras(ws)
    faixas(ws, len(df), len(df.columns))


def aba_destaques(wb: Workbook, dados: dict[str, pd.DataFrame]) -> None:
    volume = dados["1620"]
    corrente = dados["1846"]
    sazonal = dados["1621"]
    taxa = dados["5932"]

    pib_ano = volume[volume.serie == "PIB a preços de mercado"].groupby("ano").valor.mean()
    nominal = corrente[corrente.serie == "PIB a preços de mercado"].groupby("ano").valor.sum()
    va_2025 = corrente[(corrente.ano == 2025) &
                       (corrente.serie == "Valor adicionado a preços básicos")].valor.sum()
    servicos_2025 = corrente[(corrente.ano == 2025) &
                             (corrente.serie == "Serviços - total")].valor.sum()
    cagr = (pib_ano[2025] / pib_ano[2000]) ** (1 / 25) - 1
    real_2025 = pib_ano[2025] / pib_ano[2024] - 1
    sazonal_pib = sazonal[sazonal.serie == "PIB a preços de mercado"].set_index("data").valor
    qoq = sazonal_pib["2026-04-01"] / sazonal_pib["2026-01-01"] - 1
    taxa_pib = taxa[taxa.serie == "PIB a preços de mercado"].set_index("data").valor
    yoy = taxa_pib["2026-04-01"] / 100

    ws = wb.create_sheet("Destaques")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 30
    ws.append(["Indicador", "Valor", "Leitura"])
    linhas = [
        ("PIB nominal 2025", brl_milhoes(nominal[2025]), "Preços correntes"),
        ("PIB real 2025", f"{numero_br(real_2025 * 100, 2)}%", "Variação real vs. 2024"),
        ("Crescimento médio 2000–2025", f"{numero_br(cagr * 100, 2)}% a.a.", "Composto anual"),
        ("Serviços no valor adicionado", f"{numero_br(servicos_2025 / va_2025 * 100, 1)}%", "2025"),
        ("PIB 2026-T2", f"{numero_br(yoy * 100, 1)}%", "YoY (mesmo trimestre do ano anterior)"),
        ("PIB 2026-T2", f"{numero_br(qoq * 100, 2)}%", "QoQ (com ajuste sazonal)"),
    ]
    for linha in linhas:
        ws.append(list(linha))

    estilo_cabecalho(ws, 3)
    for linha in range(2, ws.max_row + 1):
        ws.cell(row=linha, column=2).alignment = Alignment(horizontal="right")
        for coluna in range(1, 4):
            ws.cell(row=linha, column=coluna).font = Font(name=FONTE, size=11)
    faixas(ws, ws.max_row - 1, 3)


def main() -> int:
    parser = argparse.ArgumentParser(description="Gera a planilha .xlsx da base tratada do PIB.")
    parser.add_argument("--dados", default=str(RAIZ / "dados"), help="Pasta dos CSV tratados")
    parser.add_argument("--saida", default=str(RAIZ / "dados" / "planilha_pib.xlsx"),
                        help="Arquivo .xlsx de saída")
    args = parser.parse_args()

    pasta = Path(args.dados)
    arquivos = {
        "1620": "fato_pib_volume.csv",
        "1621": "fato_pib_volume_sazonal.csv",
        "1846": "fato_pib_corrente.csv",
        "5932": "fato_pib_taxa.csv",
    }
    dados = {}
    for tabela, nome in arquivos.items():
        caminho = pasta / nome
        if not caminho.exists():
            print(f"[ERRO] não encontrado: {caminho}. Rode o ETL antes.", file=sys.stderr)
            return 2
        dados[tabela] = pd.read_csv(caminho)

    tempo = pd.read_csv(pasta / "dim_tempo.csv")
    serie = pd.read_csv(pasta / "dim_serie.csv")

    wb = Workbook()
    aba_leia_me(wb)
    aba_fatos(wb, dados)
    aba_dimensao(wb, "dim_tempo", tempo)
    aba_dimensao(wb, "dim_serie", serie)
    aba_destaques(wb, dados)

    destino = Path(args.saida)
    wb.save(destino)
    print(f"[OK] planilha gravada em {destino} ({destino.stat().st_size / 1024:.0f} KB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
